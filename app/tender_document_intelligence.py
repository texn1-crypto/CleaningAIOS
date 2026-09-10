from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

from .config import settings
from .models import BusinessRecord, TenderDocument


EXTRACTOR_VERSION = "tender-requirements-v1"
PRODUCT_SPEC_EXTRACTOR_VERSION = "tender-product-specification-v2"
PRODUCT_COMPARISON_VERSION = "tender-product-comparison-v1"
PRODUCT_COMPARISON_REVIEW_VERSION = "tender-product-comparison-review-v1"
SUPPORTED_SUFFIXES = {".docx", ".md", ".pdf", ".txt"}
MAX_SEGMENTS = 20_000
MAX_CANDIDATES = 500
MAX_PRODUCT_PARAMETERS = 500

_REQUIREMENT_MARKERS = (
    "обязан",
    "должен",
    "должна",
    "должны",
    "необходимо",
    "требуется",
    "требование",
    "условие",
    "срок",
    "оплата",
    "поставка",
    "оказание услуг",
    "исполнитель",
    "участник",
    "заказчик",
    "shall",
    "must",
    "required",
    "requirement",
    "deadline",
    "payment",
    "delivery",
)
_MANDATORY_MARKERS = (
    "обязан",
    "должен",
    "должна",
    "должны",
    "необходимо",
    "требуется",
    "shall",
    "must",
    "required",
)
_INJECTION_MARKERS = (
    "ignore previous instructions",
    "ignore system prompt",
    "reveal system prompt",
    "игнорируй предыдущие инструкции",
    "игнорируй системный промпт",
    "покажи системный промпт",
)
_PRODUCT_SPEC_KINDS = {
    "product_specification",
    "supplier_quote",
    "supplier_specification",
}
_TENDER_SPEC_KINDS = {
    "requirements",
    "technical_specification",
    "tender_specification",
}


class TenderDocumentExtractionError(ValueError):
    pass


class TenderDocumentReviewError(ValueError):
    pass


def _verified_storage_path(document: TenderDocument) -> Path:
    if not document.storage_path:
        raise TenderDocumentExtractionError("Tender document has not been downloaded")
    root = Path(settings.document_storage_path).resolve()
    path = Path(document.storage_path).resolve()
    try:
        path.relative_to(root)
    except ValueError as exc:
        raise TenderDocumentExtractionError(
            "Tender document is outside protected storage"
        ) from exc
    if not path.is_file():
        raise TenderDocumentExtractionError("Tender document file is unavailable")
    if path.stat().st_size > settings.max_document_bytes:
        raise TenderDocumentExtractionError("Tender document exceeds the size limit")
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise TenderDocumentExtractionError(
            "Requirement extraction supports PDF, DOCX, TXT and Markdown documents"
        )
    checksum = hashlib.sha256(path.read_bytes()).hexdigest()
    if not re.fullmatch(r"[a-fA-F0-9]{64}", document.checksum or ""):
        raise TenderDocumentExtractionError("Tender document checksum is unavailable")
    if checksum.lower() != document.checksum.lower():
        raise TenderDocumentExtractionError("Tender document checksum mismatch")
    return path


def _clean_text(value: str) -> str:
    return " ".join(value.replace("\x00", " ").split()).strip(" |•-\t")


def _segments(path: Path) -> list[tuple[str, str]]:
    suffix = path.suffix.lower()
    result: list[tuple[str, str]] = []
    try:
        if suffix == ".pdf":
            if not path.read_bytes()[:5] == b"%PDF-":
                raise TenderDocumentExtractionError("Tender document MIME does not match PDF")
            reader = PdfReader(str(path))
            for page_number, page in enumerate(reader.pages[:500], start=1):
                for paragraph_number, raw in enumerate((page.extract_text() or "").splitlines(), start=1):
                    result.append((f"page {page_number}, line {paragraph_number}", raw))
        elif suffix == ".docx":
            try:
                with ZipFile(path) as archive:
                    if "word/document.xml" not in archive.namelist():
                        raise TenderDocumentExtractionError(
                            "Tender document MIME does not match DOCX"
                        )
            except BadZipFile as exc:
                raise TenderDocumentExtractionError(
                    "Tender document MIME does not match DOCX"
                ) from exc
            source = Document(str(path))
            for paragraph_number, paragraph in enumerate(source.paragraphs, start=1):
                result.append((f"paragraph {paragraph_number}", paragraph.text))
            for table_number, table in enumerate(source.tables, start=1):
                for row_number, row in enumerate(table.rows, start=1):
                    result.append(
                        (
                            f"table {table_number}, row {row_number}",
                            " | ".join(cell.text for cell in row.cells),
                        )
                    )
        else:
            raw_bytes = path.read_bytes()
            try:
                text = raw_bytes.decode("utf-8-sig")
            except UnicodeDecodeError:
                try:
                    text = raw_bytes.decode("cp1251")
                except UnicodeDecodeError as exc:
                    raise TenderDocumentExtractionError(
                        "Tender text document encoding is unsupported"
                    ) from exc
            for line_number, line in enumerate(text.splitlines(), start=1):
                result.append((f"line {line_number}", line))
    except TenderDocumentExtractionError:
        raise
    except Exception as exc:
        raise TenderDocumentExtractionError("Tender document text extraction failed") from exc
    return result[:MAX_SEGMENTS]


def _requirement_type(text: str) -> str:
    normalized = text.lower().replace("ё", "е")
    if any(word in normalized for word in ("опыт", "лицензи", "сро", "аккредитац", "qualification", "experience")):
        return "company_qualification"
    if any(word in normalized for word in ("сертифик", "соответств", "гост", "product compliance")):
        return "product_compliance"
    if any(word in normalized for word in ("оплат", "аванс", "payment")):
        return "payment"
    if any(word in normalized for word in ("постав", "достав", "место оказания", "delivery")):
        return "delivery"
    if any(word in normalized for word in ("обеспечени", "гаранти", "security")):
        return "security"
    if any(word in normalized for word in ("срок", "deadline", "период")):
        return "deadline"
    return "scope"


def _product_source_role(document: TenderDocument) -> str:
    analysis = document.analysis or {}
    explicit = str(analysis.get("source_role") or "").strip().lower()
    if explicit in {"offered_product", "tender_requirement"}:
        return explicit
    kind = str(analysis.get("kind") or "").strip().lower()
    if kind in _PRODUCT_SPEC_KINDS:
        return "offered_product"
    if kind in _TENDER_SPEC_KINDS:
        return "tender_requirement"
    return "unknown"


def _split_product_parameter(text: str) -> tuple[str, str] | None:
    for separator in (" | ", ":", "=", " — ", " – "):
        if separator not in text:
            continue
        parameter, value = text.split(separator, 1)
        parameter = _clean_text(parameter)
        value = _clean_text(value)
        if not (2 <= len(parameter) <= 200 and 1 <= len(value) <= 500):
            continue
        if not any(character.isalpha() for character in parameter):
            continue
        if value.lower().startswith(("http://", "https://")):
            continue
        return parameter, value
    return None


def extract_product_specification_candidates(
    document: TenderDocument,
) -> tuple[dict[str, Any], bool]:
    """Extract evidence-bound parameter/value candidates without asserting a match."""

    path = _verified_storage_path(document)
    source_role = _product_source_role(document)
    previous = (document.analysis or {}).get("product_specification_extraction")
    if (
        isinstance(previous, dict)
        and previous.get("extractor_version") == PRODUCT_SPEC_EXTRACTOR_VERSION
        and previous.get("document_checksum") == document.checksum.lower()
        and previous.get("source_role") == source_role
    ):
        return previous, False

    candidates: list[dict[str, Any]] = []
    injection_locators: list[str] = []
    seen: set[tuple[str, str]] = set()
    for locator, raw in _segments(path):
        text = _clean_text(raw)
        if not text:
            continue
        normalized = text.lower().replace("ё", "е")
        if any(marker in normalized for marker in _INJECTION_MARKERS):
            injection_locators.append(locator)
            continue
        split = _split_product_parameter(text)
        if split is None:
            continue
        parameter, value = split
        normalized_parameter = parameter.lower().replace("ё", "е")
        normalized_value = value.lower().replace("ё", "е")
        fingerprint = (normalized_parameter, normalized_value)
        if fingerprint in seen:
            continue
        seen.add(fingerprint)
        code_hash = hashlib.sha256(normalized_parameter.encode("utf-8")).hexdigest()
        candidate_hash = hashlib.sha256(
            (
                f"{document.checksum.lower()}:{normalized_parameter}:"
                f"{normalized_value}:{locator}"
            ).encode("utf-8")
        ).hexdigest()
        excerpt = text[:2_000]
        candidates.append(
            {
                "candidate_hash": candidate_hash,
                "code": f"product.{code_hash[:12]}",
                "parameter": parameter,
                "value": value,
                "source_role": source_role,
                "match_status": "unknown",
                "verification_status": "needs_verification",
                "confidence": "0.0000",
                "evidence": [
                    {
                        "document_id": document.id,
                        "document_checksum": document.checksum.lower(),
                        "locator": locator,
                        "excerpt": excerpt,
                    }
                ],
                "data_class": "extracted",
                "extraction_method": "deterministic_local",
            }
        )
        if len(candidates) >= MAX_PRODUCT_PARAMETERS:
            break

    warnings: list[dict[str, Any]] = []
    if injection_locators:
        warnings.append(
            {
                "code": "prompt_injection_text_detected",
                "locators": injection_locators[:20],
                "effect": "content_isolated_no_tools_executed",
            }
        )
    if source_role == "unknown":
        warnings.append(
            {
                "code": "product_specification_role_unknown",
                "effect": "manual_document_classification_required",
            }
        )
    if not candidates:
        warnings.append(
            {
                "code": "no_product_parameter_candidates",
                "effect": "manual_review_required",
            }
        )
    result: dict[str, Any] = {
        "extractor_version": PRODUCT_SPEC_EXTRACTOR_VERSION,
        "document_checksum": document.checksum.lower(),
        "source_role": source_role,
        "status": "needs_verification",
        "candidate_count": len(candidates),
        "candidates": candidates,
        "warnings": warnings,
        "confidence_is_advisory_only": True,
        "automatic_matching_allowed": False,
        "automatic_eligibility_allowed": False,
        "automatic_submission_allowed": False,
        "external_ai_used": False,
        "content_trust": "untrusted",
        "extracted_at": datetime.now(timezone.utc).isoformat(),
    }
    document.analysis = {
        **(document.analysis or {}),
        "product_specification_extraction": result,
    }
    document.status = "analyzed"
    document.analyzed_at = datetime.now(timezone.utc).replace(tzinfo=None)
    return result, True


def review_product_specification_candidates(
    document: TenderDocument,
    *,
    document_checksum: str,
    extractor_version: str,
    decisions: list[dict[str, Any]],
    reviewed_by: str,
) -> tuple[dict[str, Any], bool]:
    extraction = (document.analysis or {}).get("product_specification_extraction")
    if not isinstance(extraction, dict):
        raise TenderDocumentReviewError(
            "Product specification must be extracted before review"
        )
    actual_checksum = str(document.checksum or "").lower()
    if document_checksum.lower() != actual_checksum:
        raise TenderDocumentReviewError("Product specification checksum is stale")
    if extraction.get("document_checksum") != actual_checksum:
        raise TenderDocumentReviewError("Extracted product specification is stale")
    if extractor_version != extraction.get("extractor_version"):
        raise TenderDocumentReviewError("Product specification extractor version is stale")
    source_role = str(extraction.get("source_role") or "unknown")
    if source_role == "unknown":
        raise TenderDocumentReviewError(
            "Product specification role must be classified before review"
        )

    candidates = extraction.get("candidates")
    if not isinstance(candidates, list) or not candidates:
        raise TenderDocumentReviewError("Product specification has no candidates to review")
    candidate_map = {
        str(candidate.get("candidate_hash")): candidate
        for candidate in candidates
        if isinstance(candidate, dict) and candidate.get("candidate_hash")
    }
    expected_hashes = set(candidate_map)
    submitted_hashes = [str(decision.get("candidate_hash") or "") for decision in decisions]
    if len(submitted_hashes) != len(set(submitted_hashes)):
        raise TenderDocumentReviewError("Product specification review has duplicate decisions")
    if set(submitted_hashes) != expected_hashes:
        raise TenderDocumentReviewError(
            "Product specification review must cover the exact extracted candidate set"
        )

    canonical_decisions: list[dict[str, Any]] = []
    reviewed_candidates: list[dict[str, Any]] = []
    for decision in sorted(decisions, key=lambda item: str(item["candidate_hash"])):
        candidate_hash = str(decision["candidate_hash"])
        candidate = candidate_map[candidate_hash]
        accepted = bool(decision.get("accepted"))
        reason = str(decision.get("reason") or "").strip()
        if not accepted and not reason:
            raise TenderDocumentReviewError(
                "Rejected product specification candidates require a reason"
            )
        parameter = str(
            decision.get("corrected_parameter") or candidate.get("parameter") or ""
        ).strip()
        value = str(
            decision.get("corrected_value") or candidate.get("value") or ""
        ).strip()
        if not parameter or not value:
            raise TenderDocumentReviewError(
                "Reviewed product specification parameter and value are required"
            )
        normalized_parameter = parameter.lower().replace("ё", "е")
        code_hash = hashlib.sha256(normalized_parameter.encode("utf-8")).hexdigest()
        canonical_decisions.append(
            {
                "candidate_hash": candidate_hash,
                "accepted": accepted,
                "corrected_parameter": parameter,
                "corrected_value": value,
                "reason": reason,
            }
        )
        reviewed_candidates.append(
            {
                "candidate_hash": candidate_hash,
                "code": f"product.{code_hash[:12]}",
                "parameter": parameter,
                "value": value,
                "source_role": source_role,
                "review_outcome": "accepted" if accepted else "rejected",
                "verification_status": "verified" if accepted else "rejected",
                "reason": reason,
                "evidence": candidate.get("evidence") or [],
                "data_class": "verified" if accepted else "rejected",
            }
        )

    canonical = {
        "document_id": document.id,
        "document_checksum": actual_checksum,
        "extractor_version": extractor_version,
        "source_role": source_role,
        "decisions": canonical_decisions,
    }
    review_hash = hashlib.sha256(
        json.dumps(
            canonical,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    previous = (document.analysis or {}).get("product_specification_review")
    if isinstance(previous, dict) and previous.get("review_hash") == review_hash:
        return previous, False

    accepted_count = sum(
        candidate["review_outcome"] == "accepted"
        for candidate in reviewed_candidates
    )
    result: dict[str, Any] = {
        "review_hash": review_hash,
        "document_checksum": actual_checksum,
        "extractor_version": extractor_version,
        "source_role": source_role,
        "status": "reviewed",
        "candidate_count": len(reviewed_candidates),
        "accepted_count": accepted_count,
        "rejected_count": len(reviewed_candidates) - accepted_count,
        "candidates": reviewed_candidates,
        "reviewed_by": reviewed_by,
        "reviewed_at": datetime.now(timezone.utc).isoformat(),
        "eligible_for_draft_comparison": accepted_count > 0,
        "automatic_matching_allowed": False,
        "automatic_eligibility_allowed": False,
        "automatic_submission_allowed": False,
    }
    document.analysis = {
        **(document.analysis or {}),
        "product_specification_review": result,
    }
    document.status = "analyzed"
    document.analyzed_at = datetime.now(timezone.utc).replace(tzinfo=None)
    return result, True


def _reviewed_product_facts(document: TenderDocument) -> tuple[dict[str, Any], list[dict[str, Any]]] | None:
    analysis = document.analysis if isinstance(document.analysis, dict) else {}
    review = analysis.get("product_specification_review")
    if not isinstance(review, dict):
        return None
    actual_checksum = str(document.checksum or "").lower()
    if review.get("status") != "reviewed":
        raise TenderDocumentReviewError("Product specification review is incomplete")
    if review.get("document_checksum") != actual_checksum:
        raise TenderDocumentReviewError("Reviewed product specification checksum is stale")
    review_hash = str(review.get("review_hash") or "").lower()
    if not re.fullmatch(r"[a-f0-9]{64}", review_hash):
        raise TenderDocumentReviewError("Product specification review hash is unavailable")
    source_role = str(review.get("source_role") or "unknown")
    if source_role not in {"tender_requirement", "offered_product"}:
        raise TenderDocumentReviewError(
            "Reviewed product specification role is unsupported"
        )
    candidates = review.get("candidates")
    if not isinstance(candidates, list):
        raise TenderDocumentReviewError("Reviewed product specification candidates are unavailable")

    verified: list[dict[str, Any]] = []
    for candidate in candidates:
        if not isinstance(candidate, dict):
            raise TenderDocumentReviewError("Reviewed product specification candidate is invalid")
        if candidate.get("verification_status") != "verified":
            continue
        code = str(candidate.get("code") or "")
        parameter = str(candidate.get("parameter") or "").strip()
        value = str(candidate.get("value") or "").strip()
        candidate_hash = str(candidate.get("candidate_hash") or "").lower()
        if (
            not re.fullmatch(r"product\.[a-f0-9]{12}", code)
            or not parameter
            or not value
            or not re.fullmatch(r"[a-f0-9]{64}", candidate_hash)
            or candidate.get("source_role") != source_role
        ):
            raise TenderDocumentReviewError(
                "Reviewed product specification candidate identity is invalid"
            )
        evidence = candidate.get("evidence")
        if not isinstance(evidence, list) or not evidence:
            raise TenderDocumentReviewError(
                "Reviewed product specification evidence is unavailable"
            )
        for reference in evidence:
            if not isinstance(reference, dict):
                raise TenderDocumentReviewError(
                    "Reviewed product specification evidence is invalid"
                )
            if (
                reference.get("document_id") != document.id
                or str(reference.get("document_checksum") or "").lower()
                != actual_checksum
                or not str(reference.get("locator") or "").strip()
                or not str(reference.get("excerpt") or "").strip()
            ):
                raise TenderDocumentReviewError(
                    "Reviewed product specification evidence is stale"
                )
        verified.append(
            {
                "candidate_hash": candidate_hash,
                "code": code,
                "parameter": parameter,
                "value": value,
                "document_id": document.id,
                "document_checksum": actual_checksum,
                "review_hash": review_hash,
                "evidence": evidence,
            }
        )
    return (
        {
            "document_id": document.id,
            "document_checksum": actual_checksum,
            "review_hash": review_hash,
            "source_role": source_role,
        },
        verified,
    )


def build_product_comparison_draft(
    tender: BusinessRecord,
    documents: list[TenderDocument],
    *,
    actor: str,
    now: datetime | None = None,
) -> tuple[dict[str, Any], bool]:
    """Persist a fail-closed draft assembled from checksum-bound reviewed facts."""

    source_bindings: list[dict[str, Any]] = []
    facts_by_role: dict[str, list[dict[str, Any]]] = {
        "tender_requirement": [],
        "offered_product": [],
    }
    role_documents: dict[str, int] = {"tender_requirement": 0, "offered_product": 0}
    for document in sorted(documents, key=lambda item: item.id):
        reviewed = _reviewed_product_facts(document)
        if reviewed is None:
            continue
        binding, facts = reviewed
        source_role = str(binding["source_role"])
        role_documents[source_role] += 1
        source_bindings.append(binding)
        facts_by_role[source_role].extend(facts)

    if not role_documents["tender_requirement"]:
        raise TenderDocumentReviewError(
            "Reviewed tender requirement product specification is required"
        )
    if not role_documents["offered_product"]:
        raise TenderDocumentReviewError(
            "Reviewed offered product specification is required"
        )
    if not facts_by_role["tender_requirement"]:
        raise TenderDocumentReviewError(
            "Reviewed tender requirement has no verified product facts"
        )
    if not facts_by_role["offered_product"]:
        raise TenderDocumentReviewError(
            "Reviewed offered product has no verified product facts"
        )

    grouped: dict[str, dict[str, list[dict[str, Any]]]] = {}
    for source_role, facts in facts_by_role.items():
        for fact in facts:
            grouped.setdefault(
                str(fact["code"]),
                {"tender_requirement": [], "offered_product": []},
            )[source_role].append(fact)

    comparisons: list[dict[str, Any]] = []
    for code in sorted(grouped):
        required = sorted(
            grouped[code]["tender_requirement"],
            key=lambda item: (item["document_id"], item["candidate_hash"]),
        )
        offered = sorted(
            grouped[code]["offered_product"],
            key=lambda item: (item["document_id"], item["candidate_hash"]),
        )
        if not required:
            signal = "missing_required"
        elif not offered:
            signal = "missing_offered"
        elif len(required) != 1 or len(offered) != 1:
            signal = "ambiguous"
        else:
            required_value = _clean_text(str(required[0]["value"])).lower().replace("ё", "е")
            offered_value = _clean_text(str(offered[0]["value"])).lower().replace("ё", "е")
            signal = "exact" if required_value == offered_value else "different"
        canonical_facts = {"required": required, "offered": offered}
        candidate_hash = hashlib.sha256(
            json.dumps(
                {"version": PRODUCT_COMPARISON_VERSION, "code": code, **canonical_facts},
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()
        comparisons.append(
            {
                "candidate_hash": candidate_hash,
                "code": code,
                "parameter": str((required or offered)[0]["parameter"]),
                "required_value": str(required[0]["value"]) if len(required) == 1 else "",
                "offered_value": str(offered[0]["value"]) if len(offered) == 1 else "",
                "comparison_signal": signal,
                "match_status": "unknown",
                "verification_status": "needs_verification",
                "required_facts": required,
                "offered_facts": offered,
                "data_class": "calculated",
            }
        )

    canonical = {
        "version": PRODUCT_COMPARISON_VERSION,
        "tender_id": tender.id,
        "source_bindings": source_bindings,
        "comparisons": comparisons,
    }
    draft_hash = hashlib.sha256(
        json.dumps(
            canonical,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    tender_data = tender.data if isinstance(tender.data, dict) else {}
    existing_history = tender_data.get("product_comparison_drafts")
    history = list(existing_history) if isinstance(existing_history, list) else []
    for existing in history:
        if isinstance(existing, dict) and existing.get("draft_hash") == draft_hash:
            return existing, False

    signal_counts = {
        signal: sum(item["comparison_signal"] == signal for item in comparisons)
        for signal in ("exact", "different", "missing_required", "missing_offered", "ambiguous")
    }
    warnings = [
        {
            "code": f"comparison_{signal}",
            "count": count,
            "effect": "manager_verification_required",
        }
        for signal, count in signal_counts.items()
        if count and signal not in {"exact", "different"}
    ]
    created_at = now or datetime.now(timezone.utc)
    if created_at.tzinfo is None:
        created_at = created_at.replace(tzinfo=timezone.utc)
    else:
        created_at = created_at.astimezone(timezone.utc)
    result: dict[str, Any] = {
        **canonical,
        "draft_hash": draft_hash,
        "status": "needs_verification",
        "comparison_count": len(comparisons),
        "paired_count": signal_counts["exact"] + signal_counts["different"],
        "comparison_signal_counts": signal_counts,
        "warnings": warnings,
        "created_by": actor,
        "created_at": created_at.isoformat(),
        "automatic_matching_allowed": False,
        "automatic_eligibility_allowed": False,
        "automatic_submission_allowed": False,
    }
    history.append(result)
    tender.data = {
        **tender_data,
        "product_comparison_drafts": history,
        "latest_product_comparison_draft_hash": draft_hash,
        "product_comparison_status": "needs_verification",
    }
    return result, True


def review_product_comparison_draft(
    tender: BusinessRecord,
    documents: list[TenderDocument],
    *,
    draft_hash: str,
    decisions: list[dict[str, Any]],
    reviewed_by: str,
    now: datetime | None = None,
) -> tuple[dict[str, Any], bool]:
    """Persist an exact-set manager review bound to the latest comparison draft."""

    tender_data = tender.data if isinstance(tender.data, dict) else {}
    if tender_data.get("latest_product_comparison_draft_hash") != draft_hash:
        raise TenderDocumentReviewError("Product comparison draft is stale")
    draft_history = tender_data.get("product_comparison_drafts")
    if not isinstance(draft_history, list):
        raise TenderDocumentReviewError("Product comparison draft is unavailable")
    draft = next(
        (
            item
            for item in draft_history
            if isinstance(item, dict) and item.get("draft_hash") == draft_hash
        ),
        None,
    )
    if draft is None or draft.get("version") != PRODUCT_COMPARISON_VERSION:
        raise TenderDocumentReviewError("Product comparison draft is unavailable")

    source_bindings = draft.get("source_bindings")
    if not isinstance(source_bindings, list) or not source_bindings:
        raise TenderDocumentReviewError("Product comparison source bindings are unavailable")
    comparisons = draft.get("comparisons")
    if not isinstance(comparisons, list) or not comparisons:
        raise TenderDocumentReviewError("Product comparison candidates are unavailable")
    draft_canonical = {
        "version": PRODUCT_COMPARISON_VERSION,
        "tender_id": tender.id,
        "source_bindings": source_bindings,
        "comparisons": comparisons,
    }
    expected_draft_hash = hashlib.sha256(
        json.dumps(
            draft_canonical,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    if draft.get("tender_id") != tender.id or expected_draft_hash != draft_hash:
        raise TenderDocumentReviewError("Product comparison draft integrity check failed")

    document_map = {document.id: document for document in documents}
    for binding in source_bindings:
        if not isinstance(binding, dict):
            raise TenderDocumentReviewError("Product comparison source binding is invalid")
        document = document_map.get(binding.get("document_id"))
        review = (
            document.analysis.get("product_specification_review")
            if document is not None and isinstance(document.analysis, dict)
            else None
        )
        if (
            document is None
            or str(document.checksum or "").lower()
            != str(binding.get("document_checksum") or "").lower()
            or not isinstance(review, dict)
            or review.get("review_hash") != binding.get("review_hash")
        ):
            raise TenderDocumentReviewError("Product comparison source binding is stale")

    candidate_map = {
        str(candidate.get("candidate_hash")): candidate
        for candidate in comparisons
        if isinstance(candidate, dict) and candidate.get("candidate_hash")
    }
    if len(candidate_map) != len(comparisons):
        raise TenderDocumentReviewError("Product comparison candidate identity is invalid")
    submitted_hashes = [str(decision.get("candidate_hash") or "") for decision in decisions]
    if len(submitted_hashes) != len(set(submitted_hashes)):
        raise TenderDocumentReviewError("Product comparison review has duplicate decisions")
    if set(submitted_hashes) != set(candidate_map):
        raise TenderDocumentReviewError(
            "Product comparison review must cover the exact candidate set"
        )

    canonical_decisions: list[dict[str, Any]] = []
    reviewed_comparisons: list[dict[str, Any]] = []
    for decision in sorted(decisions, key=lambda item: str(item["candidate_hash"])):
        candidate_hash = str(decision["candidate_hash"])
        candidate = candidate_map[candidate_hash]
        match_status = str(decision.get("match_status") or "")
        reason = str(decision.get("reason") or "").strip()
        signal = str(candidate.get("comparison_signal") or "")
        if match_status not in {"match", "mismatch", "unknown"}:
            raise TenderDocumentReviewError("Product comparison decision is unsupported")
        if signal not in {
            "exact",
            "different",
            "missing_required",
            "missing_offered",
            "ambiguous",
        }:
            raise TenderDocumentReviewError("Product comparison signal is unsupported")
        if signal in {"missing_required", "missing_offered", "ambiguous"} and match_status != "unknown":
            raise TenderDocumentReviewError(
                "Incomplete or ambiguous comparison evidence must remain unknown"
            )
        if match_status in {"mismatch", "unknown"} and not reason:
            raise TenderDocumentReviewError(
                "Mismatch and unknown product comparison decisions require a reason"
            )
        canonical_decisions.append(
            {
                "candidate_hash": candidate_hash,
                "match_status": match_status,
                "reason": reason,
            }
        )
        reviewed_comparisons.append(
            {
                **candidate,
                "match_status": match_status,
                "verification_status": (
                    "verified" if match_status in {"match", "mismatch"} else "needs_verification"
                ),
                "manager_reason": reason,
                "data_class": "verified" if match_status != "unknown" else "calculated",
            }
        )

    canonical = {
        "version": PRODUCT_COMPARISON_REVIEW_VERSION,
        "tender_id": tender.id,
        "draft_hash": draft_hash,
        "decisions": canonical_decisions,
    }
    review_hash = hashlib.sha256(
        json.dumps(
            canonical,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    existing_reviews = tender_data.get("product_comparison_reviews")
    review_history = list(existing_reviews) if isinstance(existing_reviews, list) else []
    for existing in review_history:
        if isinstance(existing, dict) and existing.get("review_hash") == review_hash:
            return existing, False

    status_counts = {
        status: sum(item["match_status"] == status for item in reviewed_comparisons)
        for status in ("match", "mismatch", "unknown")
    }
    review_status = "needs_verification" if status_counts["unknown"] else "reviewed"
    product_match = (
        "rejected"
        if status_counts["mismatch"]
        else "unverified"
        if status_counts["unknown"]
        else "verified"
    )
    reviewed_at = now or datetime.now(timezone.utc)
    if reviewed_at.tzinfo is None:
        reviewed_at = reviewed_at.replace(tzinfo=timezone.utc)
    else:
        reviewed_at = reviewed_at.astimezone(timezone.utc)
    result: dict[str, Any] = {
        **canonical,
        "review_hash": review_hash,
        "status": review_status,
        "product_match": product_match,
        "comparison_count": len(reviewed_comparisons),
        "match_status_counts": status_counts,
        "comparisons": reviewed_comparisons,
        "reviewed_by": reviewed_by,
        "reviewed_at": reviewed_at.isoformat(),
        "automatic_eligibility_allowed": False,
        "automatic_participation_allowed": False,
        "automatic_submission_allowed": False,
    }
    review_history.append(result)
    tender.data = {
        **tender_data,
        "product_comparison_reviews": review_history,
        "latest_product_comparison_review_hash": review_hash,
        "product_comparison_status": review_status,
        "product_match": product_match,
    }
    return result, True


def extract_requirement_candidates(document: TenderDocument) -> tuple[dict[str, Any], bool]:
    path = _verified_storage_path(document)
    previous = (document.analysis or {}).get("requirement_extraction")
    if (
        isinstance(previous, dict)
        and previous.get("extractor_version") == EXTRACTOR_VERSION
        and previous.get("document_checksum") == document.checksum.lower()
    ):
        return previous, False

    candidates: list[dict[str, Any]] = []
    injection_locators: list[str] = []
    seen: set[str] = set()
    for locator, raw in _segments(path):
        text = _clean_text(raw)
        if not text:
            continue
        normalized = text.lower().replace("ё", "е")
        if any(marker in normalized for marker in _INJECTION_MARKERS):
            injection_locators.append(locator)
        if len(text) < 8 or not any(marker in normalized for marker in _REQUIREMENT_MARKERS):
            continue
        excerpt = text[:2_000]
        fingerprint = hashlib.sha256(normalized.encode("utf-8")).hexdigest()
        if fingerprint in seen:
            continue
        seen.add(fingerprint)
        requirement_type = _requirement_type(normalized)
        candidates.append(
            {
                "code": f"{requirement_type}.{fingerprint[:12]}",
                "type": requirement_type,
                "description": excerpt,
                "mandatory": any(marker in normalized for marker in _MANDATORY_MARKERS),
                "status": "unknown",
                "verification_status": "needs_verification",
                "evidence": [
                    {
                        "document_id": document.id,
                        "document_checksum": document.checksum.lower(),
                        "locator": locator,
                        "excerpt": excerpt,
                    }
                ],
                "data_class": "extracted",
                "extraction_method": "deterministic_local",
            }
        )
        if len(candidates) >= MAX_CANDIDATES:
            break

    warnings: list[dict[str, Any]] = []
    if injection_locators:
        warnings.append(
            {
                "code": "prompt_injection_text_detected",
                "locators": injection_locators[:20],
                "effect": "content_isolated_no_tools_executed",
            }
        )
    if not candidates:
        warnings.append(
            {
                "code": "no_requirement_candidates",
                "effect": "manual_review_required",
            }
        )
    result: dict[str, Any] = {
        "extractor_version": EXTRACTOR_VERSION,
        "document_checksum": document.checksum.lower(),
        "status": "needs_verification",
        "candidate_count": len(candidates),
        "candidates": candidates,
        "warnings": warnings,
        "automatic_eligibility_allowed": False,
        "automatic_submission_allowed": False,
        "external_ai_used": False,
        "content_trust": "untrusted",
        "extracted_at": datetime.now(timezone.utc).isoformat(),
    }
    document.analysis = {**(document.analysis or {}), "requirement_extraction": result}
    document.status = "analyzed"
    document.analyzed_at = datetime.now(timezone.utc).replace(tzinfo=None)
    return result, True
