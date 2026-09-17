from __future__ import annotations

import csv
import hashlib
import os
import re
import tempfile
from datetime import date, datetime, timedelta, timezone
from html import escape
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors  # type: ignore[import-untyped]
from reportlab.lib.pagesizes import A4, landscape  # type: ignore[import-untyped]
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # type: ignore[import-untyped]
from reportlab.lib.units import mm  # type: ignore[import-untyped]
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # type: ignore[import-untyped]
from sqlalchemy import select
from sqlalchemy.orm import Session

from .config import settings
from .daily_owner_pack import _register_font
from .models import BusinessRecord, OutreachConsent, Suppression
from .notifications import queue_owner_notification


CONTACT_RECORD_TYPE = "contact_directory"
OUTREACH_CANDIDATE_RECORD_TYPE = "outreach_candidate"
WEEKLY_EXPORT_RECORD_TYPE = "weekly_contact_export"
EMAIL_RE = re.compile(r"^[A-Z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Z0-9.-]+\.[A-Z]{2,63}$", re.IGNORECASE)
EXPORT_KINDS = {
    "pdf": ("application/pdf", ".pdf"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xlsx"),
    "csv": ("text/csv", ".csv"),
}


def utcnow() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def normalize_directory_email(value: object) -> str:
    email = str(value or "").strip().lower().removeprefix("mailto:").split("?", 1)[0]
    return email if EMAIL_RE.fullmatch(email) else ""


def _contact_external_id(email: str) -> str:
    return f"email:{hashlib.sha256(email.encode()).hexdigest()}"


def _candidate_status(db: Session, email: str) -> str:
    if db.get(Suppression, email) is not None:
        return "suppressed"
    consent = db.get(OutreachConsent, email)
    if consent is not None and consent.status == "verified" and consent.purpose == "commercial_outreach":
        return "awaiting_owner_approval"
    return "consent_required"


def _upsert_candidate(db: Session, contact: BusinessRecord) -> tuple[BusinessRecord, bool]:
    email = str((contact.data or {}).get("email") or "")
    external_id = _contact_external_id(email)
    row = db.scalar(
        select(BusinessRecord).where(
            BusinessRecord.record_type == OUTREACH_CANDIDATE_RECORD_TYPE,
            BusinessRecord.external_id == external_id,
        )
    )
    created = row is None
    if row is None:
        row = BusinessRecord(
            record_type=OUTREACH_CANDIDATE_RECORD_TYPE,
            external_id=external_id,
            title=contact.title,
            owner="sales",
            source="contact_directory",
        )
        db.add(row)
    status = _candidate_status(db, email)
    row.status = status
    row.title = contact.title
    row.data = {
        "contact_record_id": contact.id,
        "email": email,
        "organization_name": (contact.data or {}).get("organization_name") or contact.title,
        "city": (contact.data or {}).get("city") or "",
        "inn": (contact.data or {}).get("inn") or "",
        "source_urls": list((contact.data or {}).get("source_urls") or []),
        "consent_verified": status == "awaiting_owner_approval",
        "owner_approval_required": True,
        "outbound_messages_created": False,
    }
    db.flush()
    return row, created


def upsert_contact_directory(
    db: Session,
    *,
    email: object,
    organization_name: object,
    city: object = "",
    inn: object = "",
    source_url: object,
    source_kind: object,
    baseline: bool,
    linked_record_id: int | None = None,
    observed_at: datetime | None = None,
) -> dict[str, Any]:
    """Upsert one global email identity and its source observation.

    A public address becomes a durable outreach candidate, never an outbound
    message. Consent, suppression and exact owner approval remain separate gates.
    """
    normalized = normalize_directory_email(email)
    if not normalized:
        raise ValueError("Invalid contact email")
    current = observed_at or utcnow()
    external_id = _contact_external_id(normalized)
    row = db.scalar(
        select(BusinessRecord).where(
            BusinessRecord.record_type == CONTACT_RECORD_TYPE,
            BusinessRecord.external_id == external_id,
        )
    )
    created = row is None
    old_data = dict(row.data or {}) if row is not None else {}
    observed_url = str(source_url or "").strip()[:1024]
    observed_kind = str(source_kind or "unknown").strip()[:128] or "unknown"
    observations = [dict(item) for item in old_data.get("observations") or [] if isinstance(item, dict)]
    observation = next(
        (
            item
            for item in observations
            if item.get("source_url") == observed_url and item.get("source_kind") == observed_kind
        ),
        None,
    )
    if observation is None:
        observations.append(
            {
                "source_kind": observed_kind,
                "source_url": observed_url,
                "first_seen_at": current.isoformat(),
                "last_seen_at": current.isoformat(),
            }
        )
    else:
        observation["last_seen_at"] = current.isoformat()
    links = {int(value) for value in old_data.get("linked_record_ids") or [] if str(value).isdigit()}
    if linked_record_id:
        links.add(linked_record_id)
    is_baseline = bool(old_data.get("is_baseline")) or baseline
    title = " ".join(str(organization_name or old_data.get("organization_name") or normalized).split())[:255]
    data = {
        **old_data,
        "email": normalized,
        "organization_name": title,
        "city": " ".join(str(city or old_data.get("city") or "").split())[:255],
        "inn": re.sub(r"\D", "", str(inn or old_data.get("inn") or ""))[:12],
        "is_baseline": is_baseline,
        "first_discovered_at": old_data.get("first_discovered_at") or current.isoformat(),
        "last_discovered_at": current.isoformat(),
        "source_urls": sorted({str(item.get("source_url") or "") for item in observations if item.get("source_url")}),
        "observations": observations,
        "linked_record_ids": sorted(links),
        "mailing_queue_status": "baseline_excluded" if is_baseline else _candidate_status(db, normalized),
    }
    if row is None:
        row = BusinessRecord(
            record_type=CONTACT_RECORD_TYPE,
            external_id=external_id,
            title=title,
            status="baseline" if is_baseline else "new",
            owner="research",
            source="contact_directory",
            data=data,
        )
        db.add(row)
        db.flush()
    else:
        row.title = title
        row.status = "baseline" if is_baseline else row.status
        row.data = data
        db.flush()
    candidate_created = False
    candidate_id = None
    if is_baseline:
        existing_candidate = db.scalar(
            select(BusinessRecord).where(
                BusinessRecord.record_type == OUTREACH_CANDIDATE_RECORD_TYPE,
                BusinessRecord.external_id == external_id,
            )
        )
        if existing_candidate is not None:
            existing_candidate.status = "baseline_excluded"
            existing_candidate.data = {
                **(existing_candidate.data or {}),
                "consent_verified": False,
                "baseline_excluded": True,
                "outbound_messages_created": False,
            }
            candidate_id = existing_candidate.id
    else:
        candidate, candidate_created = _upsert_candidate(db, row)
        candidate_id = candidate.id
        row.data = {**row.data, "mailing_queue_status": candidate.status}
    return {
        "contact_id": row.id,
        "created": created,
        "baseline": is_baseline,
        "candidate_id": candidate_id,
        "candidate_created": candidate_created,
        "mailing_queue_status": row.data["mailing_queue_status"],
    }


def contact_directory_summary(db: Session) -> dict[str, Any]:
    contacts = db.scalars(
        select(BusinessRecord).where(BusinessRecord.record_type == CONTACT_RECORD_TYPE)
    ).all()
    candidates = db.scalars(
        select(BusinessRecord).where(BusinessRecord.record_type == OUTREACH_CANDIDATE_RECORD_TYPE)
    ).all()
    states: dict[str, int] = {}
    for candidate in candidates:
        states[candidate.status] = states.get(candidate.status, 0) + 1
    baseline = sum(bool((row.data or {}).get("is_baseline")) for row in contacts)
    exported = sum(bool((row.data or {}).get("first_export_key")) for row in contacts)
    return {
        "total_contacts": len(contacts),
        "baseline_contacts": baseline,
        "new_contacts": len(contacts) - baseline,
        "exported_new_contacts": exported,
        "awaiting_weekly_export": sum(
            not bool((row.data or {}).get("is_baseline")) and not bool((row.data or {}).get("first_export_key"))
            for row in contacts
        ),
        "mailing_candidates": len(candidates),
        "mailing_candidate_states": dict(sorted(states.items())),
        "automatic_outbound_messages_created": 0,
        "consent_and_owner_approval_required": True,
    }


def _safe_output_path(path: Path) -> Path:
    root = Path(settings.document_storage_path).resolve()
    resolved = path.resolve()
    try:
        resolved.relative_to(root)
    except ValueError as exc:
        raise RuntimeError("Contact export path is outside document storage") from exc
    resolved.parent.mkdir(parents=True, exist_ok=True)
    return resolved


def _display_rows(contacts: list[BusinessRecord]) -> list[list[str]]:
    rows: list[list[str]] = []
    for contact in contacts:
        data = contact.data or {}
        rows.append(
            [
                str(data.get("organization_name") or contact.title),
                str(data.get("city") or ""),
                str(data.get("inn") or ""),
                str(data.get("email") or ""),
                str((data.get("source_urls") or [""])[0]),
                str(data.get("first_discovered_at") or ""),
                str(data.get("mailing_queue_status") or "consent_required"),
            ]
        )
    return rows


def _write_csv(path: Path, rows: list[list[str]]) -> str:
    path = _safe_output_path(path)
    temporary = tempfile.NamedTemporaryFile(prefix=f".{path.stem}-", suffix=".tmp.csv", dir=path.parent, delete=False)
    temporary_path = Path(temporary.name)
    temporary.close()
    try:
        with temporary_path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["Организация", "Город", "ИНН", "Email", "Источник", "Обнаружен", "Очередь рассылки"])
            writer.writerows(rows)
        os.chmod(temporary_path, 0o600)
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_xlsx(path: Path, rows: list[list[str]], *, week_start: date) -> str:
    path = _safe_output_path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Новые контакты"
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "Новые контакты УК и ТСЖ"
    sheet["A2"] = f"Неделя {week_start.isoformat()}. Исходная база и ранее выгруженные адреса исключены."
    sheet.merge_cells("A1:G1")
    sheet.merge_cells("A2:G2")
    headers = ["Организация", "Город", "ИНН", "Email", "Источник", "Обнаружен", "Очередь рассылки"]
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:G{max(3, len(rows) + 3)}"
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color="173F32")
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="657068")
    header_fill = PatternFill("solid", fgColor="174F3C")
    for cell in sheet[3]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for sheet_row in sheet.iter_rows(min_row=4):
        for cell in sheet_row:
            cell.font = Font(name="Arial", size=10, color="29332E")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in {"A": 38, "B": 22, "C": 15, "D": 34, "E": 52, "F": 22, "G": 25}.items():
        sheet.column_dimensions[column].width = width
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:G{max(3, len(rows) + 3)}"
    sheet.print_title_rows = "1:3"
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.35
    sheet.page_margins.bottom = 0.35
    temporary = tempfile.NamedTemporaryFile(prefix=f".{path.stem}-", suffix=".tmp.xlsx", dir=path.parent, delete=False)
    temporary_path = Path(temporary.name)
    temporary.close()
    try:
        workbook.save(temporary_path)
        workbook.close()
        os.chmod(temporary_path, 0o600)
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_pdf(path: Path, rows: list[list[str]], *, week_start: date, generated_at: datetime) -> str:
    path = _safe_output_path(path)
    font = _register_font()
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "WeeklyContactsTitle",
        parent=styles["Title"],
        fontName=font,
        fontSize=18,
        leading=23,
        textColor=colors.HexColor("#173f32"),
        spaceAfter=4 * mm,
    )
    body = ParagraphStyle(
        "WeeklyContactsBody",
        parent=styles["BodyText"],
        fontName=font,
        fontSize=7,
        leading=9,
        textColor=colors.HexColor("#29332e"),
    )
    small = ParagraphStyle("WeeklyContactsSmall", parent=body, fontSize=6.5, leading=8, textColor=colors.HexColor("#657068"))
    temporary = tempfile.NamedTemporaryFile(prefix=f".{path.stem}-", suffix=".tmp.pdf", dir=path.parent, delete=False)
    temporary_path = Path(temporary.name)
    temporary.close()
    document = SimpleDocTemplate(
        str(temporary_path),
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title="Новые контакты УК и ТСЖ",
        author=settings.company_name,
    )

    def paragraph(value: object, style: ParagraphStyle = body) -> Paragraph:
        text = " ".join(str(value or "").split())
        return Paragraph(escape(text[:500]), style)

    def page(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont(font, 7)
        canvas.setFillColor(colors.HexColor("#657068"))
        canvas.drawString(10 * mm, 6 * mm, "CleaningAIOS · PostgreSQL · публичные источники")
        canvas.drawRightString(landscape(A4)[0] - 10 * mm, 6 * mm, f"Страница {doc.page}")
        canvas.restoreState()

    header = ["Организация", "Город", "ИНН", "Email", "Источник", "Обнаружен"]
    table_rows = [[paragraph(value, small) for value in header]]
    table_rows.extend([[paragraph(value) for value in row[:6]] for row in rows])
    story: list[Any] = [
        paragraph("Новые контакты УК и ТСЖ", title),
        paragraph(
            f"Неделя {week_start.isoformat()} · сформировано {generated_at.isoformat()} · новых адресов: {len(rows)}. "
            "Исходные и ранее выгруженные адреса исключены.",
            small,
        ),
        Spacer(1, 3 * mm),
    ]
    if rows:
        story.append(
            Table(
                table_rows,
                repeatRows=1,
                colWidths=[54 * mm, 31 * mm, 22 * mm, 49 * mm, 91 * mm, 30 * mm],
                style=TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#174f3c")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b8c9c1")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 3),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                        ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ]
                ),
            )
        )
    else:
        story.append(paragraph("За неделю новых адресов нет.", body))
    story.extend(
        [
            Spacer(1, 3 * mm),
            paragraph(
                "Адрес в публичном источнике не является согласием на рекламу. Контакты помещены в безопасную "
                "очередь кандидатов; отправка требует проверенного согласия, suppression-проверки и точного одобрения владельца.",
                small,
            ),
        ]
    )
    try:
        document.build(story, onFirstPage=page, onLaterPages=page)
        raw = temporary_path.read_bytes()
        if not raw.startswith(b"%PDF"):
            raise RuntimeError("Weekly contact export is not a valid PDF")
        if len(raw) > settings.max_attachment_bytes:
            raise RuntimeError("Weekly contact PDF exceeds the Telegram attachment limit")
        os.chmod(temporary_path, 0o600)
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _artifact(path: Path, checksum: str, kind: str) -> dict[str, str]:
    content_type, _ = EXPORT_KINDS[kind]
    return {
        "storage_path": str(path),
        "filename": path.name,
        "sha256": checksum,
        "content_type": content_type,
    }


def verify_contact_export_artifact(report: BusinessRecord, kind: str) -> tuple[Path, dict[str, str]]:
    if kind not in EXPORT_KINDS:
        raise LookupError("Contact export format not found")
    artifact = ((report.data or {}).get("artifacts") or {}).get(kind) or {}
    path = _safe_output_path(Path(str(artifact.get("storage_path") or "")))
    content_type, suffix = EXPORT_KINDS[kind]
    if not path.is_file() or path.suffix.lower() != suffix:
        raise FileNotFoundError("Contact export artifact is unavailable")
    raw = path.read_bytes()
    checksum = str(artifact.get("sha256") or "")
    if len(checksum) != 64 or hashlib.sha256(raw).hexdigest() != checksum:
        raise RuntimeError("Contact export checksum mismatch")
    if kind == "pdf" and not raw.startswith(b"%PDF"):
        raise RuntimeError("Contact export PDF is invalid")
    return path, {
        "content_type": content_type,
        "filename": Path(str(artifact.get("filename") or path.name)).name,
    }


def weekly_export_key(current: datetime | None = None) -> tuple[str, date, datetime]:
    moment = current or datetime.now(timezone.utc)
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=timezone.utc)
    local = moment.astimezone(ZoneInfo(settings.contact_export_timezone))
    week_start = local.date() - timedelta(days=local.weekday())
    return f"week:{week_start.isoformat()}", week_start, local


def build_weekly_contact_export(
    db: Session,
    *,
    current: datetime | None = None,
    notify_owner: bool = True,
) -> dict[str, Any]:
    export_key, week_start, local_now = weekly_export_key(current)
    existing = db.scalar(
        select(BusinessRecord).where(
            BusinessRecord.record_type == WEEKLY_EXPORT_RECORD_TYPE,
            BusinessRecord.external_id == export_key,
            BusinessRecord.status == "completed",
        )
    )
    if existing is not None:
        for kind in EXPORT_KINDS:
            verify_contact_export_artifact(existing, kind)
        return {
            "status": "completed",
            "reused": True,
            "report_id": existing.id,
            "week_start": week_start.isoformat(),
            "new_contacts": int((existing.data or {}).get("contact_count") or 0),
            "artifacts": (existing.data or {}).get("artifacts") or {},
            "external_messages_sent": False,
            "evidence": [{"type": "weekly_contact_export", "report_id": existing.id, "reused": True}],
        }
    contacts = db.scalars(
        select(BusinessRecord)
        .where(BusinessRecord.record_type == CONTACT_RECORD_TYPE)
        .order_by(BusinessRecord.created_at, BusinessRecord.id)
    ).all()
    new_contacts = [
        row
        for row in contacts
        if not bool((row.data or {}).get("is_baseline")) and not bool((row.data or {}).get("first_export_key"))
    ]
    rows = _display_rows(new_contacts)
    directory = Path(settings.document_storage_path) / "reports" / "weekly-contacts" / week_start.isoformat()
    pdf_path = directory / f"new-management-contacts-{week_start.isoformat()}.pdf"
    xlsx_path = directory / f"new-management-contacts-{week_start.isoformat()}.xlsx"
    csv_path = directory / f"new-management-contacts-{week_start.isoformat()}.csv"
    artifacts = {
        "pdf": _artifact(pdf_path, _write_pdf(pdf_path, rows, week_start=week_start, generated_at=local_now), "pdf"),
        "xlsx": _artifact(xlsx_path, _write_xlsx(xlsx_path, rows, week_start=week_start), "xlsx"),
        "csv": _artifact(csv_path, _write_csv(csv_path, rows), "csv"),
    }
    report = BusinessRecord(
        record_type=WEEKLY_EXPORT_RECORD_TYPE,
        external_id=export_key,
        title=f"Новые контакты УК/ТСЖ · {week_start.isoformat()}",
        status="completed",
        owner="research",
        source="contact_directory",
        data={
            "week_start": week_start.isoformat(),
            "generated_at": local_now.isoformat(),
            "contact_count": len(new_contacts),
            "contact_record_ids": [row.id for row in new_contacts],
            "artifacts": artifacts,
            "baseline_excluded": True,
            "previous_exports_excluded": True,
        },
    )
    db.add(report)
    db.flush()
    for contact in new_contacts:
        contact.data = {
            **(contact.data or {}),
            "first_export_key": export_key,
            "first_export_report_id": report.id,
            "first_exported_at": local_now.isoformat(),
        }
    notification = None
    if notify_owner:
        pdf = artifacts["pdf"]
        notification = queue_owner_notification(
            db,
            idempotency_key=f"weekly-contact-export:{export_key}:telegram",
            channel="telegram",
            resource_type=WEEKLY_EXPORT_RECORD_TYPE,
            resource_id=str(report.id),
            subject="Новые контакты УК и ТСЖ за неделю",
            body=(
                f"Новых адресов: {len(new_contacts)}. PDF приложен; Excel и CSV сохранены в истории выпуска. "
                "Публичные адреса не отправляются без подтверждённого согласия и одобрения владельца."
            ),
            data={
                "report_id": report.id,
                "document_path": pdf["storage_path"],
                "document_filename": pdf["filename"],
                "document_sha256": pdf["sha256"],
                "document_content_type": pdf["content_type"],
            },
        )
    return {
        "status": "completed",
        "reused": False,
        "report_id": report.id,
        "week_start": week_start.isoformat(),
        "new_contacts": len(new_contacts),
        "artifacts": artifacts,
        "notification_id": notification.id if notification is not None else None,
        "external_messages_sent": False,
        "evidence": [
            {
                "type": "weekly_contact_export",
                "report_id": report.id,
                "contact_count": len(new_contacts),
                "checksums": [artifact["sha256"] for artifact in artifacts.values()],
            }
        ],
    }
